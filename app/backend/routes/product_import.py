"""Bounded Excel/CSV preview and atomic, retry-safe vendor product imports."""
import csv
import hashlib
from io import BytesIO, StringIO
from pathlib import Path
from zipfile import ZipFile, BadZipFile
from xml.etree.ElementTree import ParseError
from defusedxml.common import DefusedXmlException

from firebase_admin import firestore
from flask import Blueprint, g, jsonify, request, Response

from auth_utils import require_role
from db import SHOP_ITEMS, doc, col, db, now_iso
from routes.khata import KhataError, _khata_error
from routes.vendors import _item_patch, _vendor_profile_fields

product_import_bp = Blueprint("product_import", __name__)
product_import_bp.register_error_handler(KhataError, _khata_error)
MAX_BYTES = 2 * 1024 * 1024
MAX_ROWS = 200
FIELDS = ("name", "category", "description", "price", "stock_quantity")


def parse_products(contents, filename):
    if not contents or len(contents) > MAX_BYTES:
        raise KhataError("Choose a nonempty file up to 2 MB")
    suffix = Path(filename).suffix.lower()
    workbook = None
    try:
        if suffix == ".csv":
            text = contents.decode("utf-8-sig")
            rows = csv.reader(StringIO(text))
        elif suffix == ".xlsx":
            # Reject ZIP bombs before XML parsing. Formula cells are rejected below.
            with ZipFile(BytesIO(contents)) as archive:
                infos = archive.infolist()
                if len(infos) > 1000 or sum(info.file_size for info in infos) > 10 * 1024 * 1024:
                    raise KhataError("Workbook expands beyond the supported limit")
            from openpyxl import load_workbook
            workbook = load_workbook(BytesIO(contents), read_only=True, data_only=False, keep_links=False)
            rows = workbook.active.iter_rows(values_only=True)
        else:
            raise KhataError("Choose .xlsx or UTF-8 .csv; save older .xls files as .xlsx")
        header = next(rows, None)
        if not header or len(header) > 20:
            raise KhataError("A header row with name, price and stock_quantity is required")
        names = [str(value or "").strip().lower() for value in header]
        if len(set(names)) != len(names) or set(names) - set(FIELDS):
            raise KhataError("Use only name, category, description, price, stock_quantity headers")
        if not {"name", "price", "stock_quantity"}.issubset(names):
            raise KhataError("Required headers: name, price, stock_quantity")
        products, seen = [], set()
        for row_number, row in enumerate(rows, 2):
            if row_number > MAX_ROWS + 1:
                raise KhataError("Import at most 200 product rows per file")
            if not any(value not in (None, "") for value in row):
                continue
            if any(value not in (None, "") for value in row[len(names):]):
                raise KhataError(f"Row {row_number}: extra cells do not have headers")
            if any(isinstance(value, str) and value.startswith("=") for value in row):
                raise KhataError(f"Row {row_number}: replace formulas with values")
            values = dict(zip(names, row))
            # Excel represents an integral numeric cell as either int or float.
            stock = values.get("stock_quantity")
            if isinstance(stock, float) and stock.is_integer():
                values["stock_quantity"] = int(stock)
            for field in ("name", "category", "description"):
                values[field] = str(values.get(field) or "").strip()
            try:
                product = _item_patch(values, creating=True)
            except KhataError as exc:
                raise KhataError(f"Row {row_number}: {exc.detail}") from None
            key = (product["name"].casefold(), product["category"].casefold())
            if key in seen:
                raise KhataError(f"Row {row_number}: duplicate product name/category")
            seen.add(key)
            products.append(product)
        if not products:
            raise KhataError("No product rows found")
        return products
    except KhataError:
        raise
    except (UnicodeError, BadZipFile, csv.Error, ValueError, KeyError, TypeError, OSError, ParseError, DefusedXmlException):
        raise KhataError("File is not a readable Excel/UTF-8 CSV workbook") from None
    finally:
        if workbook:
            workbook.close()


@product_import_bp.get("/vendors/items/import/template")
@require_role("vendor")
def import_template():
    return Response("name,category,description,price,stock_quantity\nRice,Groceries,1 kg bag,60,20\n",
                    mimetype="text/csv", headers={"Content-Disposition": 'attachment; filename="products-template.csv"'})


@product_import_bp.post("/vendors/items/import")
@require_role("vendor")
def import_products():
    upload = request.files.get("file")
    if not upload:
        raise KhataError("Choose a product file")
    contents = upload.stream.read(MAX_BYTES + 1)
    products = parse_products(contents, upload.filename or "")
    if request.args.get("preview") == "true":
        return jsonify({"count": len(products), "products": products, "preview": True})
    fingerprint = hashlib.sha256(g.user_id.encode() + b"\0" + contents).hexdigest()
    import_ref = doc("product_imports", fingerprint)
    profile = _vendor_profile_fields()
    refs = [col(SHOP_ITEMS).document() for _ in products]

    @firestore.transactional
    def commit(tx):
        prior = import_ref.get(transaction=tx)
        if prior.exists:
            return prior.to_dict()["response"]
        timestamp = now_iso()
        for ref, product in zip(refs, products):
            tx.create(ref, {**profile, **product, "created_at": timestamp, "import_id": fingerprint})
        response = {"count": len(products), "item_ids": [ref.id for ref in refs], "import_id": fingerprint}
        tx.create(import_ref, {"vendor_id": g.user_id, "created_at": timestamp, "response": response})
        return response
    return jsonify(commit(db().transaction())), 201
